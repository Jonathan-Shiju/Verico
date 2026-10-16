"""The workbook a person opens after a keep.

The conversation stays in the session log. The file is a copy: a cover, the
narrative, one row per cited payment, one row per account with its entity
link, and the shorthand work log. A dismiss does not open one.
"""

import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from GraphIngest.Windows import format_ts
from GraphRead.explain import explain_transactions
from GraphRead.reader import GraphReader
from Operator.ConversationLog import ConversationLog

ROOT = Path(__file__).resolve().parent.parent
CASES = ROOT / "presentation" / "cases"

_BOLD = Font(bold=True)
_WRAP = Alignment(wrap_text=True, vertical="top")


def open_sheet(
    log: ConversationLog,
    reader: GraphReader,
    session_id: str,
    lead: str,
) -> dict:
    """Write the workbook for one kept lead and store its path on the session."""
    lead = (lead or "").strip()
    if not lead:
        return {"error": "Name the lead the workbook is for"}
    decision = reader.decision_for_lead(lead)
    if decision.get("decision") == "dismiss":
        return {"error": "A dismiss does not open a sheet"}
    if decision.get("decision") != "keep":
        return {"error": "A sheet opens after a keep"}
    explained = explain_transactions(reader, decision["txn_ids"])
    if explained.get("error"):
        return {"error": explained["error"]}
    if explained.get("missing"):
        return {"error": "The cited payments could not be loaded"}
    payments = _payments(explained, decision["txn_ids"])
    parties = _parties(explained, decision.get("accounts") or [])
    relative = f"presentation/cases/{_safe(session_id)}.xlsx"
    path = ROOT / relative
    _write(path, log.read(session_id), decision, payments, parties)
    if not log.set_sheet(session_id, relative):
        return {"error": "The session pointer is missing"}
    return {
        "session_id": session_id,
        "lead": lead,
        "sheet": relative,
        "payments": len(payments),
        "parties": len(parties),
    }


def workbook_tool(log: ConversationLog, reader: GraphReader, session_id: str):
    """A tool the model calls when the officer asks for the workbook."""

    def open_workbook(lead: str) -> dict:
        """Write the workbook for a lead the officer already kept.

        Call this when the officer asks for the workbook. lead is the lead
        already open in the conversation. Record keep for that lead first.
        A dismiss does not open a workbook.
        """
        return open_sheet(log, reader, session_id, lead)

    return open_workbook


def refresh_workbook(log: ConversationLog, reader: GraphReader, session_id: str, calls: list[str]) -> None:
    """Rewrite a workbook the model already opened, so this turn is in the narrative."""
    for call in calls:
        if not call.startswith("sheet:"):
            continue
        lead = call[len("sheet:") :]
        if lead:
            open_sheet(log, reader, session_id, lead)


def _payments(explained: dict, txn_ids: list[int]) -> list[dict]:
    by_id = {}
    for graph in explained.get("graphs") or []:
        accounts = {account["id"]: account for account in graph.get("accounts") or []}
        for transfer in graph.get("transfers") or []:
            src = accounts.get(transfer["src"]) or {}
            dst = accounts.get(transfer["dst"]) or {}
            by_id[int(transfer["txn_id"])] = {
                "txn_id": int(transfer["txn_id"]),
                "time": transfer.get("time") or "",
                "src": transfer["src"],
                "dst": transfer["dst"],
                "amount_paid": transfer.get("amount_paid"),
                "payment_currency": transfer.get("payment_currency") or "",
                "amount_received": transfer.get("amount_received"),
                "receiving_currency": transfer.get("receiving_currency") or "",
                "payment_format": transfer.get("payment_format") or "",
                "src_entity_id": src.get("entity_id") or "",
                "src_entity_name": src.get("entity_name") or "",
                "dst_entity_id": dst.get("entity_id") or "",
                "dst_entity_name": dst.get("entity_name") or "",
            }
    return [by_id[txn_id] for txn_id in txn_ids if txn_id in by_id]


def _parties(explained: dict, named: list[str]) -> list[dict]:
    found = {}
    for graph in explained.get("graphs") or []:
        for account in graph.get("accounts") or []:
            found[account["id"]] = {
                "account": account["id"],
                "entity_id": account.get("entity_id") or "",
                "entity_name": account.get("entity_name") or "",
                "on_keep": "yes" if account["id"] in named else "",
            }
    return [found[account] for account in sorted(found)]


def _write(path: Path, turns: list[dict], decision: dict, payments: list[dict], parties: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    cover = book.active
    cover.title = "Cover"
    recorded = decision.get("recorded_ts")
    rows = [
        ("Session", path.stem),
        ("Lead", decision.get("lead") or ""),
        ("Decision", decision.get("decision") or ""),
        ("Reason", decision.get("reason") or ""),
        ("Recorded", format_ts(int(recorded)) if recorded else ""),
        ("Named accounts", ", ".join(decision.get("accounts") or [])),
        ("Payments", len(payments)),
        ("Parties", len(parties)),
    ]
    for index, (label, value) in enumerate(rows, 1):
        cover.cell(index, 1, label).font = _BOLD
        cell = cover.cell(index, 2, value)
        cell.alignment = _WRAP
    cover.column_dimensions["A"].width = 18
    cover.column_dimensions["B"].width = 72

    _table(
        book.create_sheet("Narrative"),
        ["Time", "Speaker", "Text"],
        [[_when(turn), turn.get("role") or "", turn.get("text") or ""] for turn in turns],
        wrap_column=3,
    )
    _table(
        book.create_sheet("Payments"),
        [
            "Transaction",
            "Time",
            "From account",
            "To account",
            "Amount paid",
            "Paid currency",
            "Amount received",
            "Received currency",
            "Format",
            "From entity",
            "From entity name",
            "To entity",
            "To entity name",
        ],
        [
            [
                payment["txn_id"],
                payment["time"],
                payment["src"],
                payment["dst"],
                payment["amount_paid"],
                payment["payment_currency"],
                payment["amount_received"],
                payment["receiving_currency"],
                payment["payment_format"],
                payment["src_entity_id"],
                payment["src_entity_name"],
                payment["dst_entity_id"],
                payment["dst_entity_name"],
            ]
            for payment in payments
        ],
    )
    _table(
        book.create_sheet("Parties"),
        ["Account", "Entity id", "Entity name", "Named on keep"],
        [
            [party["account"], party["entity_id"], party["entity_name"], party["on_keep"]]
            for party in parties
        ],
    )
    _table(
        book.create_sheet("Work log"),
        ["Time", "Call"],
        [
            [_when(turn), call]
            for turn in turns
            for call in (turn.get("calls") or [])
        ],
    )
    book.save(path)


def _table(sheet, headers: list[str], rows: list[list], wrap_column: int = 0) -> None:
    for column, header in enumerate(headers, 1):
        sheet.cell(1, column, header).font = _BOLD
    for row_index, row in enumerate(rows, 2):
        for column, value in enumerate(row, 1):
            cell = sheet.cell(row_index, column, "" if value is None else value)
            if column == wrap_column:
                cell.alignment = _WRAP
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in range(1, len(headers) + 1):
        width = len(headers[column - 1])
        for row in rows:
            width = max(width, min(len(str(row[column - 1] if row[column - 1] is not None else "")), 42))
        sheet.column_dimensions[get_column_letter(column)].width = width + 2


def _when(turn: dict) -> str:
    raw = turn.get("ts")
    if not raw:
        return ""
    return format_ts(int(raw))


def _safe(session_id: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]", "_", session_id)
